from __future__ import annotations

from io import BytesIO

from app.services.ingestion.models import NormalizedDocument, NormalizedSection, SourceFile
from app.services.ingestion.parsers.base import DocumentParser


class DocxParser(DocumentParser):
    def supports(self, source: SourceFile) -> bool:
        return source.extension == ".docx"

    def parse(self, content: bytes, source: SourceFile) -> NormalizedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX parser dependency missing. Install python-docx.") from exc

        document = Document(BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text and paragraph.text.strip()]
        text = "\n\n".join(paragraphs)
        section = NormalizedSection(section_id="body", order=0, heading="Document Body", text=text)
        warnings = [] if text else ["No extractable paragraph text found in DOCX."]
        return NormalizedDocument(
            source=source,
            sections=[section] if text else [],
            raw_text=text,
            metadata={"paragraphs": len(paragraphs)},
            parse_warnings=warnings,
        )
